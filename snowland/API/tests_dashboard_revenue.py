from datetime import datetime, time, timezone as utc_timezone
from unittest.mock import patch
from io import BytesIO
import csv
import io

from django.contrib.auth.models import User
from django.test import TestCase
from django.utils import timezone
from rest_framework.test import APIClient

from Client.models import Client
from Control.models import UserProfile
from Coursekit.models import CourseCategory, CourseTemplate, CourseType
from Resorts.models import Campus, Resorts
from booking.models import Booking, Payment, Reservation, ReservationGroup, SharedCourseClass


class DashboardRevenueTests(TestCase):
    def setUp(self):
        self.tenant = Client.objects.create(name='SnowLand', internal_code='snowland')
        self.campus = Campus.objects.create(client=self.tenant, name='測試校區', code='qa')
        self.resort = Resorts.objects.create(client=self.tenant, name='qa', display_name='測試雪場')
        self.campus.resorts.add(self.resort)
        self.client = APIClient()
        self.client.force_authenticate(User.objects.create_superuser('root', 'root@example.test', 'test-password'))

    def order(self, *, status='completed', payment='paid', amount=1001, bookings=2):
        group = ReservationGroup.objects.create(client=self.tenant, campus=self.campus, marketing_source='Google')
        reservation = Reservation.objects.create(group=group, resort=self.resort, status=status, course_fee=amount)
        for _ in range(bookings):
            Booking.objects.create(reservation=reservation, date=timezone.localdate(), start_time=time(9),
                                   end_time=time(11), course_type='單板', course_name='上午班')
        if payment:
            Payment.objects.create(reservation_group=group, status=payment)
        return group

    def dashboard(self, params=None):
        response = self.client.get('/api/admin/snowland/dashboard/', params or {})
        self.assertEqual(response.status_code, 200, response.data)
        return response.data['data']

    def test_only_paid_fully_completed_orders_contribute_revenue(self):
        self.order()
        for status in ('created', 'cancelled', 'deleted', 'manually_assigned'):
            self.order(status=status, amount=9000)
        for payment in ('unpaid', 'pending', 'expired', 'refunded', None):
            self.order(payment=payment, amount=9000)
        self.order(bookings=0, amount=9000)
        data = self.dashboard()
        self.assertEqual(data['stats']['orders'], 11)
        self.assertEqual(data['stats']['revenue'], 1001)
        self.assertEqual(data['stats']['completed_orders'], 1)
        self.assertEqual(data['stats']['avg_order_value'], 1001)
        self.assertEqual(data['marketing_sources'][0]['revenue'], 1001)
        self.assertEqual(data['campus_summary'][0]['revenue'], 1001)
        self.assertEqual(data['top_items'], [{'id': 1, 'name': '上午班', 'quantity': 2, 'revenue': 1001}])

    def test_partial_completion_does_not_recognize_entire_order_early(self):
        group = self.order()
        Reservation.objects.create(group=group, resort=self.resort, status='manually_assigned', course_fee=2000)
        self.assertEqual(self.dashboard()['stats']['revenue'], 0)

    def test_deleted_reservation_is_not_revenue_or_a_completion_blocker(self):
        group = self.order()
        Reservation.objects.create(group=group, resort=self.resort, status='deleted', course_fee=9000)
        self.assertEqual(self.dashboard()['stats']['revenue'], 1001)

    def test_latest_payment_controls_recognition_not_any_old_paid_attempt(self):
        group = self.order()
        Payment.objects.create(reservation_group=group, status='refunded')
        self.assertEqual(self.dashboard()['stats']['revenue'], 0)
        Payment.objects.create(reservation_group=group, status='paid')
        self.assertEqual(self.dashboard()['stats']['revenue'], 1001)

    def test_discounts_and_average_use_completed_orders_not_all_created_orders(self):
        group = self.order(amount=2000)
        reservation = group.reservations.get()
        reservation.discount_amount = 500
        reservation.save()
        self.order(amount=2500)
        self.order(status='created', amount=9000)
        data = self.dashboard()
        self.assertEqual(data['stats']['revenue'], 4000)
        self.assertEqual(data['stats']['avg_order_value'], 2000)
        self.assertEqual(sum(row['revenue'] for row in data['top_items']), 4000)

    def test_no_completed_orders_returns_zero_average(self):
        self.order(status='created')
        data = self.dashboard()
        self.assertEqual(data['stats']['revenue'], 0)
        self.assertEqual(data['stats']['avg_order_value'], 0)

    def test_custom_dates_include_end_day_and_use_taipei_midnight(self):
        for timestamp in ((15, 59), (16, 0), (23, 59)):
            group = self.order()
            ReservationGroup.objects.filter(pk=group.pk).update(
                created_at=datetime(2026, 10, 15, *timestamp, tzinfo=utc_timezone.utc))
        following = self.order()
        ReservationGroup.objects.filter(pk=following.pk).update(
            created_at=datetime(2026, 10, 16, 16, 0, tzinfo=utc_timezone.utc))
        data = self.dashboard({'period':'custom', 'start_date':'2026-10-16', 'end_date':'2026-10-16'})
        self.assertEqual(data['stats']['orders'], 2)
        self.assertEqual(data['stats']['revenue'], 2002)

    def test_invalid_filters_are_rejected_instead_of_silently_showing_month(self):
        for params in (
            {'period':'unknown'}, {'date_basis':'unknown'}, {'period':'custom'},
            {'period':'custom','start_date':'2026-10-17','end_date':'2026-10-16'},
            {'period':'custom','start_date':'2026-02-30','end_date':'2026-10-16'},
        ):
            with self.subTest(params=params):
                response = self.client.get('/api/admin/snowland/dashboard/', params)
                self.assertEqual(response.status_code, 400, response.data)

    def test_completion_date_filter_finds_old_orders_completed_in_selected_period(self):
        group = self.order()
        ReservationGroup.objects.filter(pk=group.pk).update(
            created_at=datetime(2025, 1, 1, tzinfo=utc_timezone.utc))
        group.reservations.get().bookings.update(date='2026-10-17')
        params = {'period':'custom','start_date':'2026-10-17','end_date':'2026-10-17'}
        self.assertEqual(self.dashboard(params)['stats']['orders'], 0)
        completed = self.dashboard({**params, 'date_basis':'completed'})
        self.assertEqual(completed['stats']['orders'], 1)
        self.assertEqual(completed['stats']['revenue'], 1001)
        unfinished = self.order(status='created')
        unfinished.reservations.get().bookings.update(date='2026-10-17')
        self.assertEqual(self.dashboard({**params,'date_basis':'completed'})['stats']['orders'], 1)

    def test_quarter_does_not_include_previous_or_next_quarter(self):
        for month in (9, 10, 12):
            group = self.order()
            ReservationGroup.objects.filter(pk=group.pk).update(
                created_at=datetime(2026, month, 15, tzinfo=utc_timezone.utc))
        group = self.order()
        ReservationGroup.objects.filter(pk=group.pk).update(
            created_at=datetime(2027, 1, 1, tzinfo=utc_timezone.utc))
        with patch('API.admin_views.timezone.now', return_value=datetime(2026, 10, 17, tzinfo=utc_timezone.utc)):
            self.assertEqual(self.dashboard({'period':'quarter'})['stats']['orders'], 2)

    def test_resort_filter_counts_only_matching_reservations_not_entire_mixed_order(self):
        other = Resorts.objects.create(client=self.tenant, name='other-qa', display_name='另一雪場')
        self.campus.resorts.add(other)
        group = self.order()
        reservation = Reservation.objects.create(group=group, resort=other, status='completed', course_fee=3000, number_of_people=3)
        Booking.objects.create(reservation=reservation,date=timezone.localdate(),start_time=time(9),end_time=time(12),course_name='另一課程')
        data = self.dashboard({'resort_id':str(other.pk)})
        self.assertEqual(data['stats']['orders'], 1)
        self.assertEqual(data['stats']['revenue'], 3000)
        self.assertEqual(data['stats']['student_visits'], 3)
        self.assertEqual(data['stats']['teaching_hours'], 3)
        self.assertEqual(data['top_items'][0]['name'], '另一課程')
        self.assertEqual(data['recent_orders'][0]['amount'], 3000)
        self.assertEqual({row['id'] for row in data['resort_options']}, {self.resort.pk,other.pk})

    def test_resort_filter_never_promotes_a_partially_completed_order_to_completed(self):
        other = Resorts.objects.create(client=self.tenant,name='later-qa',display_name='未完課雪場')
        self.campus.resorts.add(other)
        group = self.order()
        Reservation.objects.create(group=group,resort=other,status='created',course_fee=9000)
        data = self.dashboard({'resort_id':str(self.resort.pk)})
        self.assertEqual(data['stats']['orders'],1)
        self.assertEqual(data['stats']['revenue'],0)

    def test_resort_filter_and_export_cannot_escape_campus_or_company(self):
        self.order()
        other_campus = Campus.objects.create(client=self.tenant,name='另一校區',code='other')
        other_resort = Resorts.objects.create(client=self.tenant,name='outside-qa',display_name='不該出現雪場')
        other_campus.resorts.add(other_resort)
        foreign = Client.objects.create(name='外公司',internal_code='outside')
        foreign_resort = Resorts.objects.create(client=foreign,name='foreign-qa',display_name='外公司雪場')
        user = User.objects.create_user('analytics-manager','qa-manager@example.test','test-password')
        profile = UserProfile.objects.get(user=user)
        profile.is_manager=True
        profile.manager_client=self.tenant
        profile.role='campus_manager'
        profile.admin_permissions=['analytics']
        profile.save()
        profile.campuses.add(self.campus)
        self.client.force_authenticate(User.objects.get(pk=user.pk))
        self.assertEqual({row['id'] for row in self.dashboard()['resort_options']},{self.resort.pk})
        for resort in (other_resort,foreign_resort):
            for export in ('','csv','xlsx'):
                response=self.client.get('/api/admin/snowland/dashboard/',{'resort_id':str(resort.pk),'export':export})
                self.assertEqual(response.status_code,403,response.content[:200])
        # A no-longer-assigned resort remains selectable for authorized historical orders.
        self.campus.resorts.clear()
        self.assertEqual(self.dashboard({'resort_id':str(self.resort.pk)})['stats']['orders'],1)

    def test_invalid_resort_or_export_format_is_not_silently_ignored(self):
        for params in ({'resort_id':'oops'},{'resort_id':'-1'},{'export':'pdf'}):
            response=self.client.get('/api/admin/snowland/dashboard/',params)
            self.assertEqual(response.status_code,400,response.content[:200])

    def test_completed_student_visits_hours_and_payment_statuses_are_explicit(self):
        group=self.order()
        reservation=group.reservations.get()
        reservation.number_of_people=3
        reservation.save()
        self.order(payment='unpaid')
        self.order(status='manually_assigned',payment='pending')
        data=self.dashboard()
        self.assertEqual(data['stats']['student_visits'],6)
        self.assertEqual(data['stats']['teaching_hours'],4)
        self.assertEqual(data['payment_statuses'],[
            {'status':'paid','name':'已付款','orders':1},
            {'status':'pending','name':'待確認付款','orders':1},
            {'status':'unpaid','name':'未付款','orders':1},
            {'status':'expired','name':'付款時間已過期','orders':0},
            {'status':'refunded','name':'已退款','orders':0},
        ])

    def test_shared_class_hours_count_once_but_each_order_keeps_its_students(self):
        category=CourseCategory.objects.create(client=self.tenant,name='單板')
        kind=CourseType.objects.create(category=category,name='團體')
        template=CourseTemplate.objects.create(course_type=kind,name='團體班',duration_hours=2)
        shared=SharedCourseClass.objects.create(client=self.tenant,campus=self.campus,resort=self.resort,
            template=template,date=timezone.localdate(),start_time=time(9),end_time=time(11),ability_level='no_exp',status='ready')
        for _ in range(3):
            group=self.order(bookings=1)
            reservation=group.reservations.get()
            reservation.shared_course_class=shared
            reservation.number_of_people=1
            reservation.save()
        data=self.dashboard()
        self.assertEqual(data['stats']['student_visits'],3)
        self.assertEqual(data['stats']['teaching_hours'],2)
        self.assertEqual(data['stats']['revenue'],3003)

    def test_export_requires_analytics_permission_and_authentication(self):
        user=User.objects.create_user('no-analytics','no-analytics@example.test','test-password')
        profile=UserProfile.objects.get(user=user)
        profile.is_manager=True
        profile.manager_client=self.tenant
        profile.role='assistant'
        profile.admin_permissions=['orders']
        profile.save()
        self.client.force_authenticate(User.objects.get(pk=user.pk))
        self.assertEqual(self.client.get('/api/admin/snowland/dashboard/',{'export':'xlsx'}).status_code,403)
        self.client.force_authenticate(None)
        self.assertIn(self.client.get('/api/admin/snowland/dashboard/',{'export':'csv'}).status_code,(401,403))

    def test_csv_and_excel_export_same_filtered_summary_without_member_contacts(self):
        group=self.order()
        group.marketing_source='=HYPERLINK("https://example.test","bad")'
        group.booking_contact={'email':'never-export@example.test','phone':'never-export-phone'}
        group.save()
        params={'resort_id':str(self.resort.pk),'export':'csv'}
        response=self.client.get('/api/admin/snowland/dashboard/',params)
        self.assertEqual(response.status_code,200)
        self.assertIn('attachment;',response['Content-Disposition'])
        content=response.content.decode('utf-8-sig')
        self.assertNotIn('never-export',content)
        rows=list(csv.reader(io.StringIO(content)))
        self.assertIn(['已完成訂單營業額', '1001'],rows)
        self.assertTrue(any(row and row[0].startswith("'=HYPERLINK") for row in rows))
        excel=self.client.get('/api/admin/snowland/dashboard/',{**params,'export':'xlsx'})
        self.assertEqual(excel.status_code,200)
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(excel.content))
        self.assertTrue(book.active['B6'].alignment.wrap_text)
        self.assertGreaterEqual(book.active.row_dimensions[6].height, 90)
        cells=[cell for row in book.active for cell in row]
        self.assertFalse(any(cell.data_type=='f' for cell in cells))
        self.assertTrue(any(cell.value==group.marketing_source and cell.data_type=='s' for cell in cells))
        self.assertFalse(any('never-export' in str(cell.value) for cell in cells))

    def test_export_handles_pasted_control_characters_without_formula_execution(self):
        group=self.order()
        group.marketing_source='\x0b=HYPERLINK("https://example.test","QA")\x00'
        group.save()
        for file_format in ('xlsx','csv'):
            with self.subTest(file_format=file_format):
                response=self.client.get('/api/admin/snowland/dashboard/',{'export':file_format})
                self.assertEqual(response.status_code,200)
                if file_format=='xlsx':
                    from openpyxl import load_workbook
                    cells=[cell for row in load_workbook(BytesIO(response.content)).active for cell in row]
                    self.assertFalse(any(cell.data_type=='f' for cell in cells))
                    self.assertTrue(any(cell.value=='=HYPERLINK("https://example.test","QA")' for cell in cells))
                else:
                    content=response.content.decode('utf-8-sig')
                    self.assertNotIn('\x00',content);self.assertNotIn('\x0b',content)
                    rows=list(csv.reader(io.StringIO(content)))
                    self.assertTrue(any(row and row[0]=='\'=HYPERLINK("https://example.test","QA")' for row in rows))
        group.refresh_from_db()
        self.assertIn('\x00',group.marketing_source)
