from collections import defaultdict
from datetime import datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.contrib.auth.models import User
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from Client.models import Client
from Coach.models import Coach
from Coursekit.models import CourseCategory, CourseTemplate
from Resorts.models import Resorts
from booking.models import Booking, MemberDetail, Payment, Reservation, ReservationGroup


RESORT_ALIASES = {
    '星野': ('星野', 'tomamu'),
    '富良野': ('富良野', 'furano'),
    '留壽都': ('留壽都', 'rusutsu'),
    '神居': ('神居', 'kamui'),
    '喜樂樂': ('喜樂樂', 'kiroro'),
    '札幌國際': ('札幌國際', 'sapporo koksai'),
    '手稻': ('手稻', 'teine'),
    '佐幌': ('佐幌', 'sahoro'),
    '夕張': ('夕張', 'yubari'),
}

COURSE_TIMES = {
    '全天': (time(9), time(15), 6),
    '半天_早上': (time(9), time(12), 3),
    '半天_下午': (time(13), time(16), 3),
    '半天_上下下午均可依排課': (time(9), time(12), 3),
    '野雪嚮導': (time(9), time(15), 6),
}


def clean_text(value):
    return str(value or '').strip()


def clean_order_number(value):
    if value in (None, ''):
        return ''
    try:
        number = Decimal(str(value))
        return str(int(number)) if number == number.to_integral_value() else format(number, 'f')
    except (InvalidOperation, ValueError):
        return clean_text(value)


def age_range(value):
    try:
        age = int(float(value))
    except (TypeError, ValueError):
        return '19-24y'
    for upper, label in ((6, '4-6y'), (12, '7-12y'), (18, '13-18y'), (24, '19-24y'), (35, '25-35y'), (45, '36-45y'), (55, '46-55y')):
        if age <= upper:
            return label
    return '56y以上'


class Command(BaseCommand):
    help = 'Validate or import the customer 26/27 order workbook. Dry-run is the default.'

    def add_arguments(self, parser):
        parser.add_argument('file')
        parser.add_argument('--client-code', default='snowland')
        parser.add_argument('--commit', action='store_true')
        parser.add_argument('--skip-invalid', action='store_true')

    def handle(self, *args, **options):
        source = Path(options['file'])
        if not source.is_file():
            raise CommandError(f'找不到檔案：{source}')
        client = Client.objects.filter(internal_code=options['client_code']).first()
        if not client:
            raise CommandError(f'找不到客戶：{options["client_code"]}')

        workbook = load_workbook(source, read_only=True, data_only=True)
        if '訂單統計表' not in workbook.sheetnames:
            raise CommandError('找不到工作表：訂單統計表')
        sheet = workbook['訂單統計表']
        grouped = defaultdict(list)
        for row_number, cells in enumerate(sheet.iter_rows(min_row=3), start=3):
            row = {
                get_column_letter(index): cell.value
                for index, cell in enumerate(cells, start=1)
            }
            if not any(row.get(key) not in (None, '') for key in ('P', 'V', 'W', 'AE', 'AB')):
                continue
            key = clean_order_number(row.get('DA') or row.get('CZ') or row.get('P'))
            grouped[key].append((row_number, row))
        workbook.close()

        invalid = {}
        warnings = []
        resolved = {}
        for key, rows in grouped.items():
            issues = []
            email = clean_text(rows[0][1].get('W')).lower()
            if not email or '@' not in email:
                issues.append('缺少有效 Email，無法綁定會員')
            for row_number, row in rows:
                resort_name = clean_text(row.get('AB'))
                course_name = clean_text(row.get('AD'))
                resort = self._resolve_resort(client, resort_name)
                if not resort:
                    issues.append(f'第 {row_number} 列找不到雪場設定：{resort_name or "空白"}')
                if course_name not in COURSE_TIMES:
                    issues.append(f'第 {row_number} 列課程時段無法判斷：{course_name or "空白"}')
                if resort:
                    for discipline, people in self._disciplines(row):
                        template = self._resolve_template(client, resort, discipline, COURSE_TIMES.get(course_name, (None, None, 0))[2])
                        if people and not template:
                            issues.append(f'第 {row_number} 列缺少 {discipline} {COURSE_TIMES.get(course_name, (None, None, 0))[2]} 小時課程模板')
                if email.endswith('.con'):
                    warnings.append(f'{key}：Email 網域為 .con，請人工確認')
            if issues:
                invalid[key] = sorted(set(issues))
            else:
                resolved[key] = rows

        self.stdout.write(f'讀取 {len(grouped)} 筆訂單群組、{sum(map(len, grouped.values()))} 列課程')
        self.stdout.write(self.style.SUCCESS(f'可匯入：{len(resolved)} 筆'))
        self.stdout.write(self.style.WARNING(f'需補資料／設定：{len(invalid)} 筆'))
        for key, issues in list(invalid.items())[:20]:
            self.stdout.write(self.style.WARNING(f'- {key}: {"；".join(issues)}'))
        if len(invalid) > 20:
            self.stdout.write(self.style.WARNING(f'- 另有 {len(invalid) - 20} 筆未逐筆顯示'))
        for warning in sorted(set(warnings)):
            self.stdout.write(self.style.WARNING(f'- 提醒：{warning}'))

        if not options['commit']:
            self.stdout.write('DRY RUN：未寫入資料庫。確認後加上 --commit。')
            return
        if invalid and not options['skip_invalid']:
            raise CommandError('存在無法匯入的訂單。補齊資料，或使用 --skip-invalid 僅匯入可處理訂單。')

        created = skipped = 0
        for key, rows in resolved.items():
            if ReservationGroup.objects.filter(client=client, order_number=key).exists():
                skipped += 1
                continue
            with transaction.atomic():
                self._import_group(client, key, rows)
                created += 1
        self.stdout.write(self.style.SUCCESS(f'完成：新增 {created} 筆，已存在略過 {skipped} 筆。'))

    def _resolve_resort(self, client, source_name):
        aliases = RESORT_ALIASES.get(source_name)
        if not aliases:
            return None
        query = Resorts.objects.filter(client=client)
        for alias in aliases:
            result = query.filter(display_name__icontains=alias).first() or query.filter(name__icontains=alias).first()
            if result:
                return result
        return None

    def _resolve_template(self, client, resort, discipline, duration):
        keyword = '單板' if discipline == 'snowboard' else '雙板'
        return CourseTemplate.objects.filter(
            course_type__category__client=client,
            course_type__category__name__icontains=keyword,
            resorts=resort,
            duration_hours=duration,
            is_active=True,
        ).order_by('display_order', 'id').first()

    @staticmethod
    def _disciplines(row):
        values = []
        for discipline, column in (('snowboard', 'AI'), ('ski', 'AJ')):
            text = clean_text(row.get(column))
            try:
                people = int(text.rsplit('-', 1)[-1].replace('人', '')) if text else 0
            except ValueError:
                people = 0
            if people:
                values.append((discipline, people))
        if not values:
            total = int(float(row.get('AG') or 0))
            if total:
                values.append(('snowboard', total))
        return values

    def _import_group(self, client, key, rows):
        first = rows[0][1]
        email = clean_text(first.get('W')).lower()
        user = User.objects.filter(email__iexact=email).first()
        if not user:
            username = f'snowland_{key}'[:150]
            suffix = 1
            base = username
            while User.objects.filter(username=username).exists():
                suffix += 1
                username = f'{base[:140]}_{suffix}'
            user = User.objects.create_user(username=username, email=email)
            user.set_unusable_password()
            user.save(update_fields=['password'])
        full_name = clean_text(first.get('V'))
        if full_name and not user.get_full_name():
            user.first_name = full_name
            user.save(update_fields=['first_name'])

        first_resort = self._resolve_resort(client, clean_text(first.get('AB')))
        campus = first_resort.campuses.filter(client=client, is_active=True).first()
        created_at = first.get('R') or first.get('CY')
        group = ReservationGroup.objects.create(
            client=client,
            campus=campus,
            order_number=key,
            name=full_name or key,
            marketing_source=clean_text(first.get('S')),
            marketing_source_detail='既有訂單匯入',
            source_ip=clean_text(first.get('CX')) or None,
            user=user,
        )
        if isinstance(created_at, datetime):
            ReservationGroup.objects.filter(pk=group.pk).update(created_at=timezone.make_aware(created_at) if timezone.is_naive(created_at) else created_at)

        total_paid = 0
        payment_method = 'TT'
        for row_number, row in rows:
            amount = self._money(row.get('DF')) + self._money(row.get('DG')) + self._money(row.get('DH')) + self._money(row.get('DI')) + self._money(row.get('DJ'))
            total_paid += max(amount, 0)
            if self._money(row.get('DG')) or self._money(row.get('DH')):
                payment_method = 'newebpay'
            disciplines = self._disciplines(row)
            people_total = sum(value for _, value in disciplines) or 1
            for discipline, people in disciplines:
                resort = self._resolve_resort(client, clean_text(row.get('AB')))
                start_time, end_time, duration = COURSE_TIMES[clean_text(row.get('AD'))]
                template = self._resolve_template(client, resort, discipline, duration)
                category = CourseCategory.objects.filter(client=client, name__icontains='單板' if discipline == 'snowboard' else '雙板').first()
                course_type = template.course_type if template else (category.types.first() if category else None)
                requested_coach = clean_text(row.get('AL') or row.get('AK'))
                coach = None
                if requested_coach and '無指定' not in requested_coach:
                    coach = Coach.objects.filter(client=client, name__iexact=requested_coach).first()
                share = max(round(max(amount, 0) * people / people_total), 0)
                reservation = Reservation.objects.create(
                    group=group,
                    resort=resort,
                    course_type=course_type,
                    course_template=template,
                    is_preferred_coach=bool(coach and clean_text(row.get('AK')) and '無指定' not in clean_text(row.get('AK'))),
                    preferred_coach=coach,
                    language='zh' if '中' in clean_text(row.get('AM')) else 'en',
                    status='manually_assigned' if coach else 'manual_assignment_needed',
                    number_of_people=people,
                    need_equipment=clean_text(row.get('AN')).startswith('是'),
                    course_fee=share,
                )
                start_date = row.get('AE')
                days = max(int(float(row.get('AF') or 1)), 1)
                if isinstance(start_date, datetime):
                    start_date = start_date.date()
                for offset in range(days):
                    Booking.objects.create(
                        reservation=reservation,
                        course_type='單板' if discipline == 'snowboard' else '雙板',
                        course_name=clean_text(row.get('AD')),
                        date=start_date + timedelta(days=offset),
                        start_time=start_time,
                        end_time=end_time,
                        is_scheduled=bool(coach),
                    )
                self._create_members(reservation, user, row, discipline, people)

        Payment.objects.create(
            reservation_group=group,
            user=user,
            status='paid' if total_paid > 0 else 'pending',
            payment_method=payment_method,
            DataJSON={
                'import_source': '2627_訂單清單表_匯入系統',
                'source_order_number': key,
                'amount': total_paid,
                'contact': {
                    'name': full_name,
                    'email': email,
                    'phone': clean_text(first.get('X')),
                    'messenger_type': clean_text(first.get('Y')),
                    'messenger_id': clean_text(first.get('Z')),
                },
            },
        )

    @staticmethod
    def _money(value):
        try:
            return int(round(float(value or 0)))
        except (TypeError, ValueError):
            return 0

    def _create_members(self, reservation, user, row, discipline, people):
        skill_columns = {
            'snowboard': (('AR', 'AS', 'AT', 'AU'), ('BB', 'BC', 'BD', 'BE'), ('BL', 'BM', 'BN', 'BO'), ('BV', 'BW', 'BX', 'BY'), ('CF', 'CG', 'CH', 'CI'), ('CP', 'CQ', 'CR', 'CS')),
            'ski': (('AV', 'AW', 'AX', 'AY'), ('BF', 'BG', 'BH', 'BI'), ('BP', 'BQ', 'BR', 'BS'), ('BZ', 'CA', 'CB', 'CC'), ('CJ', 'CK', 'CL', 'CM'), ('CT', 'CU', 'CV', 'CW')),
        }
        age_columns = ('AQ', 'BA', 'BK', 'BU', 'CE', 'CO')
        valid = dict(MemberDetail.SNOWBOARD_CHOICES if discipline == 'snowboard' else MemberDetail.SKI_CHOICES)
        for index in range(people):
            skills = []
            if index < len(skill_columns[discipline]):
                for column in skill_columns[discipline][index]:
                    value = clean_text(row.get(column))
                    if value in valid and value not in skills:
                        skills.append(value)
            values = {
                'reservation': reservation,
                'user': user if index == 0 else None,
                'guardian': None if index == 0 else user,
                'age_range': age_range(row.get(age_columns[index]) if index < len(age_columns) else None),
            }
            values['snowboard_skills' if discipline == 'snowboard' else 'ski_skills'] = skills
            MemberDetail.objects.create(**values)
